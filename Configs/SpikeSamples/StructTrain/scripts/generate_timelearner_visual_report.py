#!/usr/bin/env python3
"""Build one Word visual report for one or more immutable TimeLearner run bundles.

The report only reads run artifacts. It never modifies configuration or trainer state.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


SCRIPT_DIR = Path(__file__).resolve().parent
METRICS_PATH = SCRIPT_DIR / "selectivity_metrics.py"
_metrics_spec = importlib.util.spec_from_file_location("timelearner_selectivity_metrics", METRICS_PATH)
if _metrics_spec is None or _metrics_spec.loader is None:
    raise RuntimeError(f"cannot load canonical metrics module: {METRICS_PATH}")
_metrics = importlib.util.module_from_spec(_metrics_spec)
_metrics_spec.loader.exec_module(_metrics)


def read_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def sha_text(value: Any) -> str:
    return str(value or "")


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def row_int(row: dict[str, str], key: str, fallback: int = -1) -> int:
    try:
        return int(float(row.get(key, "")))
    except (TypeError, ValueError):
        return fallback


def row_float(row: dict[str, str], key: str) -> float | None:
    try:
        value = float(row.get(key, ""))
        return value if math.isfinite(value) else None
    except (TypeError, ValueError):
        return None


def parse_times(row: dict[str, str], field: str) -> list[float]:
    raw = (row.get(field) or "").strip()
    out: list[float] = []
    for part in raw.split(";"):
        try:
            value = float(part)
            if math.isfinite(value):
                out.append(value)
        except ValueError:
            continue
    return out


def stimulus_times_s(row: dict[str, str]) -> list[float]:
    raw = (row.get("stim_times") or "").strip()
    if raw:
        return parse_times(row, "stim_times")
    count = row_int(row, "stim_count", 0)
    times = [0.0]
    acc = 0.0
    for index in range(max(0, count - 1)):
        acc += row_float(row, f"isi{index}") or 0.0
        times.append(acc)
    return times[:count]


def config_entries(run_root: Path, provenance: dict[str, Any]) -> list[dict[str, str]]:
    entries: list[dict[str, str]] = []
    nested_provenance = provenance.get("run_provenance") or {}
    before = provenance.get("config_sha256_before") or nested_provenance.get("config_sha256_before") or {}
    input_manifest = read_json(run_root / "inputs_manifest.json")
    copied_hashes = {
        (str(item.get("side")), str(item.get("rel"))): str(item.get("sha256"))
        for item in input_manifest.get("copied", [])
        if isinstance(item, dict)
    }
    for side, subdir in (("Train", "prelaunch-inputs/Train"), ("Test", "effective-inputs/Test/post-gate")):
        source_kind = "exact prelaunch C++ snapshot" if side == "Train" else "effective post-gate Test snapshot"
        if side == "Train" and not (run_root / subdir).is_dir():
            candidates = list(run_root.glob("runtime-input-overrides/**/Train"))
            if candidates:
                subdir = candidates[0].relative_to(run_root).as_posix()
                source_kind = "matrix input override; final prelaunch XML not preserved"
        if side == "Test" and not (run_root / subdir).is_dir():
            candidates = list(run_root.glob("runtime-input-overrides/**/Test"))
            if candidates:
                subdir = candidates[0].relative_to(run_root).as_posix()
                source_kind = "matrix Test input override; effective gate XML not preserved"
        for filename in ("Model_00.xml", "Parameters_00.xml", "Project.ini", "Interface.xml"):
            copied = run_root / subdir / filename
            if copied.is_file():
                recorded_hash = (
                    provenance.get("prelaunch_train_config_snapshot", {}).get("files", {}).get(filename, {}).get("sha256")
                    if side == "Train"
                    else provenance.get("test_config_snapshots", {}).get("post_gate", {}).get("files", {}).get(filename, {}).get("sha256")
                )
                entries.append(
                    {
                        "side": side,
                        "filename": filename,
                        "path": str(copied),
                        "source": source_kind,
                        "sha256": sha_text(recorded_hash)
                        or copied_hashes.get((side, filename))
                        or sha_text(before.get(f"{side}/{filename}"))
                        or hashlib.sha256(copied.read_bytes()).hexdigest(),
                    }
                )
            elif before.get(f"{side}/{filename}"):
                entries.append(
                    {
                        "side": side,
                        "filename": filename,
                        "path": "not copied into this historical run bundle",
                        "source": "path unavailable in run bundle",
                        "sha256": str(before[f"{side}/{filename}"]),
                    }
                )
    return entries


def locate_run_artifact(run_root: Path, relative: str, work_root: Path | None = None) -> Path | None:
    direct = run_root / relative
    if direct.is_file():
        return direct
    if work_root is not None:
        candidate = work_root / relative
        if candidate.is_file():
            return candidate
    return None


def downsample_trace(trace_path: Path, width: int) -> dict[int, dict[str, Any]]:
    """Stream trace twice; preserve per-pixel min/max without retaining raw rows."""
    ranges: dict[int, list[float]] = {}
    counts: dict[int, int] = {}
    with trace_path.open("r", encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            trial = row_int(row, "trial")
            time_rel = row_float(row, "time_rel_s")
            if trial < 0 or time_rel is None:
                continue
            if trial not in ranges:
                ranges[trial] = [time_rel, time_rel]
                counts[trial] = 0
            ranges[trial][0] = min(ranges[trial][0], time_rel)
            ranges[trial][1] = max(ranges[trial][1], time_rel)
            counts[trial] += 1

    buckets: dict[int, dict[str, Any]] = {
        trial: {
            "min_s": limits[0], "max_s": limits[1], "count": counts[trial],
            "width": width, "soma": {}, "ltz": {},
        }
        for trial, limits in ranges.items()
    }
    with trace_path.open("r", encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            trial = row_int(row, "trial")
            current = buckets.get(trial)
            time_rel = row_float(row, "time_rel_s")
            if current is None or time_rel is None:
                continue
            span = current["max_s"] - current["min_s"]
            column = 0 if span <= 0 else min(width - 1, max(0, int((time_rel - current["min_s"]) / span * (width - 1))))
            for key, field in (("soma", "soma_sum"), ("ltz", "ltz_potential")):
                value = row_float(row, field)
                if value is None:
                    continue
                bucket = current[key].setdefault(column, [value, value])
                bucket[0] = min(bucket[0], value)
                bucket[1] = max(bucket[1], value)
    return buckets


def artifact_runtime(node_override: str | None, modules_override: str | None) -> tuple[Path, Path]:
    dependency_root = Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies" / "node"
    node_name = "node.exe" if os.name == "nt" else "node"
    node_path = Path(node_override) if node_override else dependency_root / "bin" / node_name
    modules_path = Path(modules_override) if modules_override else dependency_root / "node_modules"
    if not node_path.is_file():
        found = shutil.which("node")
        if found:
            node_path = Path(found)
    if not node_path.is_file():
        raise FileNotFoundError("Node runtime not found; set CODEX_NODE_RUNTIME or --node-runtime")
    if not (modules_path / "@oai" / "artifact-tool").exists():
        raise FileNotFoundError("@oai/artifact-tool not found; set CODEX_NODE_MODULES or --artifact-node-modules")
    return node_path, modules_path


def link_node_modules(link_path: Path, modules_path: Path) -> None:
    if link_path.exists():
        return
    if os.name == "nt":
        command = f"New-Item -ItemType Junction -Path '{link_path}' -Target '{modules_path}' | Out-Null"
        subprocess.run(["powershell", "-NoProfile", "-Command", command], check=True, capture_output=True, text=True)
    else:
        link_path.symlink_to(modules_path, target_is_directory=True)


def xlsx_trial_payload(row: dict[str, str], trace: dict[str, Any] | None) -> dict[str, Any]:
    stim = stimulus_times_s(row)
    spikes = parse_times(row, "neuron_spike_times")
    start_s = trace["min_s"] if trace else 0.0
    end_s = trace["max_s"] if trace else max(stim + spikes + [0.0]) + 0.005
    if end_s <= start_s:
        end_s = start_s + 0.001
    points: list[dict[str, Any]] = []
    if trace:
        columns = set(trace["soma"]) | set(trace["ltz"])
        span = trace["max_s"] - trace["min_s"]
        denominator = max(1, int(trace.get("width", 1350)) - 1)
        for column in sorted(columns):
            soma = trace["soma"].get(column)
            ltz = trace["ltz"].get(column)
            points.append(
                {
                    "time_ms": (trace["min_s"] + (float(column) / denominator) * span) * 1000.0,
                    "soma_min": soma[0] if soma else None,
                    "soma_max": soma[1] if soma else None,
                    "ltz_min": ltz[0] if ltz else None,
                    "ltz_max": ltz[1] if ltz else None,
                }
            )
    return {
        "trial": row_int(row, "trial", 0),
        "start_ms": start_s * 1000.0,
        "end_ms": end_s * 1000.0,
        "input_ms": [value * 1000.0 for value in stim],
        "output_ms": [value * 1000.0 for value in spikes],
        "trace": points,
    }


def render_chart_images(
    rows: list[dict[str, str]],
    trace_data: dict[int, dict[str, Any]],
    run_id: str,
    temp_dir: Path,
    node_runtime: str | None,
    artifact_node_modules: str | None,
) -> dict[int, Path]:
    node_path, modules_path = artifact_runtime(node_runtime, artifact_node_modules)
    safe_run = re.sub(r"[^A-Za-z0-9_-]+", "_", run_id)
    runtime_dir = temp_dir / f"xlsx-chart-{safe_run}"
    runtime_dir.mkdir(parents=True, exist_ok=True)
    link_node_modules(runtime_dir / "node_modules", modules_path)
    helper_src = SCRIPT_DIR / "generate_timelearner_chart_workbook.mjs"
    helper_copy = runtime_dir / helper_src.name
    shutil.copy2(helper_src, helper_copy)
    output_dir = runtime_dir / "rendered"
    payload = {
        "trials": [
            xlsx_trial_payload(row, trace_data.get(row_int(row, "trial")))
            for row in rows
        ]
    }
    payload_path = runtime_dir / "chart_payload.json"
    payload_path.write_text(json.dumps(payload, ensure_ascii=False, allow_nan=False), encoding="utf-8")
    subprocess.run([str(node_path), str(helper_copy), str(payload_path), str(output_dir)], check=True)
    output: dict[int, Path] = {}
    for row in rows:
        trial = row_int(row, "trial", 0)
        image_path = output_dir / f"trial_{trial}.png"
        if not image_path.is_file():
            raise FileNotFoundError(f"XLSX renderer did not create a chart for trial {trial}: {image_path}")
        with Image.open(image_path) as image:
            left = min(40, max(0, image.width - 1))
            top = min(20, max(0, image.height - 1))
            image.crop((left, top, image.width, image.height)).save(image_path)
        output[trial] = image_path
    return output


def add_metadata_table(
    document: Document, items: list[tuple[str, str]], *, font_size: float = 9.0
) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.columns[0].width = Inches(1.7)
    table.columns[1].width = Inches(8.8)
    for key, value in items:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value or "—"
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for cell in cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
        for run in cells[0].paragraphs[0].runs:
            run.bold = True


def add_run(
    document: Document,
    run_root: Path,
    allow_missing_trace: bool,
    temp_dir: Path,
    node_runtime: str | None,
    artifact_node_modules: str | None,
) -> int:
    provenance_path = run_root / "provenance.json"
    provenance = read_json(provenance_path)
    if not provenance:
        raise ValueError(f"missing or invalid provenance: {provenance_path}")
    work_value = provenance.get("work_root")
    work_root = Path(work_value) if work_value else None
    csv_path = locate_run_artifact(run_root, "Test/results.csv", work_root)
    if csv_path is None:
        raise FileNotFoundError(f"Test/results.csv missing for {run_root}")
    rows = read_csv_rows(csv_path)
    metrics = _metrics.classify(rows)
    trace_path = locate_run_artifact(run_root, "Test/test_trace.csv", work_root)
    if trace_path is None and not allow_missing_trace:
        raise FileNotFoundError(
            f"test soma trace missing for {run_root}; use --allow-missing-trace to produce an events-only historical report"
        )
    trace_data = downsample_trace(trace_path, 1350) if trace_path else None

    matrix_result = read_json(run_root / "matrix_result.json")
    result = provenance.get("result") or matrix_result.get("result") or matrix_result
    factors = matrix_result.get("matrix_factors") or {}
    input_manifest = read_json(run_root / "inputs_manifest.json")
    case = str(provenance.get("case") or result.get("case") or run_root.name)
    run_id = str(provenance.get("run_id") or run_root.name)
    exact_started = provenance.get("experiment_started_utc") or provenance.get("started_utc")
    started = exact_started or input_manifest.get("utc") or "unknown (timestamp not present in run bundle)"
    started_source = "runner prelaunch timestamp" if exact_started else "input manifest setup timestamp; exact Console start unavailable"
    if isinstance(started, str) and re.fullmatch(r"\d{8}T\d{6}Z?", started):
        started = datetime.strptime(started.rstrip("Z"), "%Y%m%dT%H%M%S").replace(tzinfo=timezone.utc).isoformat()
    exact_test_started = provenance.get("test_gate_started_utc")
    test_started = exact_test_started or (
        datetime.fromtimestamp(csv_path.stat().st_mtime, timezone.utc).isoformat()
        if csv_path.is_file() else "unknown"
    )
    test_time_source = "runner timestamp before invoking Test gate script" if exact_test_started else "Test/results.csv filesystem timestamp"
    exact_finished = provenance.get("experiment_finished_utc")
    finished = exact_finished or (
        datetime.fromtimestamp(provenance_path.stat().st_mtime, timezone.utc).isoformat()
        if provenance_path.is_file() else "unknown"
    )
    finish_source = "C++ run provenance" if exact_finished else "provenance.json filesystem timestamp"
    converged = provenance.get("training_convergence") or result.get("training_convergence") or "unknown"
    config_rows = config_entries(run_root, provenance)
    commands = provenance.get("script_execution") or {}
    train_exec = commands.get("training") or {}
    gate_exec = commands.get("test_gate") or {}
    gate_metrics = gate_exec.get("metrics") or result.get("quality_metrics") or {}
    scripts = provenance.get("runner_scripts") or []
    worker_log = run_root / "worker.log"
    worker_lines = (
        worker_log.read_text(encoding="utf-8", errors="replace").splitlines()
        if worker_log.is_file() else []
    )
    run_provenance = provenance.get("run_provenance") or {}
    effective_work_root = (
        provenance.get("work_root")
        or result.get("run_dir")
        or run_provenance.get("work_root")
        or ""
    )
    work_train_ini = (
        str(PurePosixPath(str(effective_work_root)) / "Train" / "Project.ini")
        if str(effective_work_root).startswith("/")
        else str(Path(effective_work_root) / "Train" / "Project.ini")
    )
    train_command = train_exec.get("command") or (
        [
            str(provenance.get("nm") or run_provenance.get("nm") or "NeuroModelerConsole"),
            "-c", work_train_ini,
            "-s", "-t", str(provenance.get("train_t_effective") or result.get("train_t_effective") or factors.get("train_t") or ""),
            "-x", "-S",
        ] if effective_work_root else []
    )
    gate_command_lines = [line.strip() for line in worker_lines if line.lstrip().startswith("GATE ")]
    gate_commands = gate_exec.get("commands") or gate_command_lines
    if not started or started == "unknown":
        started = next((line.split("archive=", 1)[1].split(" work=", 1)[0] for line in worker_lines if line.startswith("=== CASE ") and "archive=" in line and " work=" in line), started)

    document.add_heading(f"Опыт {case}", level=1)
    trainer_kind = (
        result.get("trainer")
        or (provenance.get("experiment_spec") or {}).get("trainer")
        or factors.get("trainer")
        or result.get("kind")
        or provenance.get("trainer")
        or ""
    )
    trainer_names = {
        "branch": "NNeuronTimeLearnerBranch (Branch)",
        "asym": "NNeuronTimeLearner (Classic)",
        "phase6": "NNeuronTimeLearner (Classic / Phase6 config)",
        "classic": "NNeuronTimeLearner (Classic)",
    }
    trainer = trainer_names.get(str(trainer_kind).lower(), str(trainer_kind) or "см. Train Model_00.xml")
    mode = result.get("normalization_mode", factors.get("normalization_mode", provenance.get("normalization_mode", "не записан")))
    psi = result.get("presynaptic_inhibition", factors.get("presynaptic_inhibition", provenance.get("presynaptic_inhibition", "не записано")))
    acc = f"{metrics.get('acc_legacy', 0)}/{metrics.get('n', 0)} ({(100 * metrics.get('acc_legacy', 0) / metrics.get('n', 1)):.1f}%)" if metrics.get("n") else "недоступна"
    acc_strict = f"{metrics.get('acc_strict', 0)}/{metrics.get('n', 0)} ({(100 * metrics.get('acc_strict', 0) / metrics.get('n', 1)):.1f}%)" if metrics.get("n") else "недоступна"
    detected = "да" if metrics.get("target_hit_strict") else "нет/не подтверждено"
    script_acc = gate_metrics.get("acc_legacy", gate_metrics.get("acc"))
    script_n = gate_metrics.get("n", metrics.get("n", 0))
    script_acc_text = (
        f"{script_acc}/{script_n} ({(100 * float(script_acc) / float(script_n)):.1f}%)"
        if script_acc is not None and script_n else "не записана gate-скриптом"
    )
    script_acc_strict = gate_metrics.get("acc_strict")
    script_acc_strict_text = (
        f"{script_acc_strict}/{script_n} ({(100 * float(script_acc_strict) / float(script_n)):.1f}%)"
        if script_acc_strict is not None and script_n else "не записана gate-скриптом"
    )
    summary = [
        ("Run ID", run_id),
        ("Тренер / режим", f"{trainer}; NormalizationMode={mode}; пресинаптическое торможение={psi}; паттерн={factors.get('duration_ms', (provenance.get('experiment_spec') or {}).get('span_ms', 'не указан'))} ms; вариант={factors.get('pattern_family', '')}"),
        ("Эксперимент", f"начало: {started} ({started_source}); Test: {test_started} ({test_time_source}); завершение: {finished} ({finish_source})"),
        ("Сходимость C++", str(converged)),
        ("Обнаружение цели", detected),
        ("Точность (legacy)", acc),
        ("Точность (strict)", acc_strict),
        ("Точность из результата gate-скрипта", f"legacy={script_acc_text}; strict={script_acc_strict_text}"),
        ("fires / matches", f"{metrics.get('fires', '')} / {metrics.get('matches', '')}"),
        ("Ложные срабатывания", f"legacy={metrics.get('fa_legacy', metrics.get('fa', 0))}; strict={metrics.get('fa_strict', 0)}; late FP={metrics.get('late_fp', 0)}"),
        ("Качество формы ответа", f"{metrics.get('response_quality', 'unknown')}; ok_audit={metrics.get('ok_audit', 0)}; версия метрик={metrics.get('metrics_version', 'unknown')}"),
        ("Run bundle", str(run_root)),
        ("Исходный набор конфигов", str(input_manifest.get("archive_root") or provenance.get("archive_root") or "не записан")),
        ("Console SHA-256", str(provenance.get("binary_sha256_at_launch") or provenance.get("binary_sha256") or "не записан")),
        ("Git SHA", f"root={provenance.get('root_git_at_launch') or provenance.get('root_git')}; PulseLib={(provenance.get('gitlinks_at_launch') or provenance.get('gitlinks') or {}).get('Libraries/Nmsdk-PulseLib')}; Bin={(provenance.get('gitlinks_at_launch') or provenance.get('gitlinks') or {}).get('Bin')}"),
        ("Трасса потенциала", f"{trace_path} ({sum(v['count'] for v in trace_data.values())} samples)" if trace_path and trace_data else "не сохранена в исходном запуске; график ниже показывает только события, это не осциллограмма"),
        ("Train C++ команда", " ".join(str(v) for v in train_command) or "см. worker.log; Console command не сохранена"),
        ("Train status / rc", f"{train_exec.get('status', converged)} / {train_exec.get('return_code', result.get('child_rc', 'n/a'))}"),
        ("Test gate команды", "\n".join(" ".join(map(str, cmd)) if isinstance(cmd, (list, tuple)) else str(cmd) for cmd in gate_commands) or "не записаны"),
        ("Test gate результат", f"evaluated={gate_exec.get('evaluated', result.get('gate_evaluated', 'n/a'))}; rc={gate_exec.get('return_code', result.get('gate_rc', 'n/a'))}; ok={gate_exec.get('ok', result.get('gate_ok', 'n/a'))}; error={gate_exec.get('fail_reason', '')}"),
        ("Метрики и вывод gate-скрипта", f"metrics_line: {gate_exec.get('metrics_line', result.get('metrics', 'не сохранена'))}\nТочность: legacy={script_acc_text}; strict={script_acc_strict_text}; fires={gate_metrics.get('fires', 'не записано')}; matches={gate_metrics.get('matches', 'не записано')}; target_hit_strict={gate_metrics.get('target_hit_strict', 'не записано')}; false positives strict={gate_metrics.get('fa_strict', gate_metrics.get('fa_strict_v2', 'не записано'))}; ok_audit={gate_metrics.get('ok_audit', 'не записано')}; response_quality={gate_metrics.get('response_quality', 'не записано')}; metrics_version={gate_metrics.get('metrics_version', 'не записана')}"),
    ]
    add_metadata_table(document, summary)

    document.add_heading("Конфиги, фактически использованные в запуске", level=2)
    config_items = []
    if config_rows:
        for item in config_rows:
            config_items.append((f"{item['side']}/{item['filename']}", f"{item['source']}\n{item['path']}\nSHA-256: {item['sha256']}"))
    else:
        config_items.append(("Конфиги", "Конфиг-снимки и хеши отсутствуют в историческом bundle."))
    add_metadata_table(document, config_items, font_size=8.0)

    document.add_heading("Скрипты и происхождение результата", level=2)
    script_lines = [
        f"{entry.get('role')}: {entry.get('path')} | SHA-256={entry.get('sha256')}"
        for entry in scripts
    ]
    launcher_hashes = matrix_result.get("matrix_launcher_sha256")
    if isinstance(launcher_hashes, dict):
        script_lines.extend(
            f"matrix launcher {path} | SHA-256={digest}"
            for path, digest in launcher_hashes.items()
        )
    elif launcher_hashes:
        script_lines.append(f"matrix launcher SHA-256: {launcher_hashes}")
    script_lines.append(f"Orchestrator argv: {' '.join(map(str, commands.get('orchestrator', {}).get('argv', [])))}")
    script_lines.extend(
        f"log {entry.get('path')} | SHA-256={entry.get('sha256')}"
        for entry in commands.get("logs", [])
    )
    if worker_log.is_file():
        script_lines.append(f"worker.log: {worker_log} | SHA-256={hashlib.sha256(worker_log.read_bytes()).hexdigest()}")
        script_lines.extend(
            line.strip() for line in worker_lines
            if line.lstrip().startswith(("gate rc", "post-flag:", "  C++_TRAINING_FAILURE", "  ABORT"))
        )
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run("\n".join(script_lines) if script_lines else "Сведения о запущенных скриптах не сохранены.").font.size = Pt(8)

    document.add_heading("Реакции по тестовым примерам", level=2)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headings = ("trial", "class", "fires", "spike times, s", "error", "input times, ms")
    for cell, heading in zip(table.rows[0].cells, headings):
        cell.text = heading
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        stim = ", ".join(f"{value * 1000:.2f}" for value in stimulus_times_s(row))
        cells = table.add_row().cells
        values = (
            row.get("trial", ""), row.get("target_class", ""), row.get("neuron_fired", ""),
            row.get("neuron_spike_times", ""), row.get("error_class", ""), stim,
        )
        for cell, value in zip(cells, values):
            cell.text = value

    if rows:
        document.add_page_break()
    chart_images = render_chart_images(
        rows,
        trace_data or {},
        f"{case}_{run_id}",
        temp_dir,
        node_runtime,
        artifact_node_modules,
    )
    has_waveforms = any(data.get("soma") or data.get("ltz") for data in (trace_data or {}).values())
    for index, row in enumerate(rows):
        if index > 0 and (has_waveforms or index % 2 == 0):
            document.add_page_break()
        trial = row_int(row, "trial", index)
        image_path = chart_images[trial]
        if not trace_data or not (trace_data.get(trial, {}).get("soma") or trace_data.get(trial, {}).get("ltz")):
            note = document.add_paragraph(
                "Осциллограмма выхода сомы не записана в исходном запуске; показаны фактические входные импульсы и импульсы нейрона из results.csv."
            )
            note.paragraph_format.space_after = Pt(4)
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image_path), width=Inches(7.7))
    return len(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-root", action="append", required=True, type=Path, help="immutable run directory; repeat for a series")
    parser.add_argument("--out", required=True, type=Path, help="output .docx path")
    parser.add_argument("--allow-missing-trace", action="store_true", help="for historical runs; mark soma waveform unavailable")
    parser.add_argument("--node-runtime", default=os.environ.get("CODEX_NODE_RUNTIME"), help="Node runtime used for XLSX chart generation")
    parser.add_argument("--artifact-node-modules", default=os.environ.get("CODEX_NODE_MODULES"), help="workspace node_modules containing @oai/artifact-tool")
    args = parser.parse_args()
    args.out.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(35, 40, 45)
    for style_name, font_size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(font_size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        ppr = style._element.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)

    document.add_heading("Отчёт по визуальному анализу TimeLearner", 0)
    document.add_paragraph(f"Создано UTC: {datetime.now(timezone.utc).isoformat()}")
    document.add_paragraph(
        "Отчёт объединяет неизменяемые run bundles. Сходимость C++-обучения и качество классификации приведены отдельными показателями. "
        "Точность пересчитана каноническим selectivity_metrics.py из сохранённого Test/results.csv."
    )

    with tempfile.TemporaryDirectory(prefix="timelearner-report-") as temp_name:
        temp_dir = Path(temp_name)
        total_trials = 0
        for index, run_root in enumerate(args.run_root):
            if index:
                document.add_page_break()
            total_trials += add_run(
                document,
                run_root.resolve(),
                args.allow_missing_trace,
                temp_dir,
                args.node_runtime,
                args.artifact_node_modules,
            )
        document.add_paragraph(f"\nСводка: запусков — {len(args.run_root)}; тестовых примеров — {total_trials}.")
        document.save(args.out)
    print(f"WROTE {args.out} runs={len(args.run_root)} trials={total_trials}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
